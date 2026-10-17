"""
DeLaval Scorecard - monthly data preparation.

Reads every raw export in "Source Files", applies the cleaning rules the
Power BI report used to apply in Power Query, keeps only the report month and
writes ONE workbook (Data/Scorecard_Data.xlsx) with one sheet per model table.
The Power BI dashboard reads only that workbook (its "DataFile" parameter
must point to the same file as DEFAULT_OUTPUT below).

Usage (from the "DeLaval Scorecard" folder):
    python scripts/prepare_scorecard_data.py                  # month taken from the source files
    python scripts/prepare_scorecard_data.py --month 2026-09  # a specific month

Source files are recognised by their content (column headers / sheet names),
not by their file name, so "October ROTA.xlsx" or "Oct 2026 Team FCR.xlsx"
work without any change. Keep exactly one file of each kind in the folder.
"""

from __future__ import annotations

import argparse
import calendar
import csv
import datetime as dt
import json
import re
import sys
import warnings
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Font

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

# =========================================================================== #
#  PATHS - replace these with the real folders on your computer.               #
#  Keep the r before the quotes (it makes the backslashes work).               #
# =========================================================================== #

# Folder where you put the monthly source files (CSV and Excel exports).
DEFAULT_SOURCE = Path(r"C:\DeLaval Scorecard\Source Files")                      # <-- CHANGE

# The single Excel file this script creates. Power BI reads this file.
DEFAULT_OUTPUT = Path(r"C:\DeLaval Scorecard\Data\Scorecard_Data.xlsx")          # <-- CHANGE

# Folder with agent_aliases.csv and settings.json.
CONFIG_DIR = Path(r"C:\DeLaval Scorecard\config")                                 # <-- CHANGE

# =========================================================================== #

LOG: list[tuple[str, str]] = []


def log(level: str, message: str) -> None:
    LOG.append((level, message))
    print(f"[{level}] {message}")


# --------------------------------------------------------------------------- #
# Configuration and agent names
# --------------------------------------------------------------------------- #
def load_settings() -> dict:
    with open(CONFIG_DIR / "settings.json", encoding="utf-8") as f:
        return json.load(f)


def load_aliases() -> dict[str, str]:
    aliases = {}
    with open(CONFIG_DIR / "agent_aliases.csv", encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            alias, agent = (row.get("Alias") or "").strip(), (row.get("Agent") or "").strip()
            if alias and agent:
                aliases[_key(alias)] = agent
    return aliases


def _key(name: str) -> str:
    return re.sub(r"\s+", " ", str(name)).strip().lower()


def proper(text):
    """Power Query Text.Proper equivalent (first letter of each word upper-case)."""
    if text is None or (isinstance(text, float) and pd.isna(text)):
        return text
    return str(text).title()


class Names:
    """Maps any spelling of an agent's name (EXT suffix, upper/lower case,
    old surnames) to the single display name used in the Agents table."""

    def __init__(self, aliases: dict[str, str]):
        self.aliases = aliases
        self.roster: dict[str, str] = {}

    def set_roster(self, agents: list[str]) -> None:
        self.roster = {_key(a): a for a in agents}

    def canon(self, raw):
        if raw is None or (isinstance(raw, float) and pd.isna(raw)):
            return None
        name = re.sub(r"\s+", " ", str(raw)).strip()
        if not name:
            return ""
        name = re.sub(r"(?<=[a-z])EXT$", "", name).strip()  # "Amr AbdelwahabEXT"
        k = _key(name)
        if k in self.aliases:
            return self.aliases[k]
        if k in self.roster:
            return self.roster[k]
        return proper(name)

    def in_roster(self, name) -> bool:
        return name is not None and _key(name) in self.roster


# --------------------------------------------------------------------------- #
# Source file discovery
# --------------------------------------------------------------------------- #
CSV_KINDS = {
    "genesys_calls": lambda h: {"Conversation ID", "Total Queue", "Users - Interacted"} <= h,
    "genesys_status": lambda h: {"Agent Name", "Primary Status", "Start Time"} <= h,
    "csat": lambda h: {"instance", "string_value", "instance.taken_on"} <= h,
    "chat_timed_out": lambda h: {"rejected_by", "work_item"} <= h,
    "chat_closed": lambda h: {"number", "opened_at", "opened_for", "type", "state", "assigned_to"} <= h,
    "incidents": lambda h: {"number", "opened_at", "caller_id", "reassignment_count"} <= h,
    "tasks": lambda h: {"number", "opened_at", "request_item", "reassignment_count"} <= h and "caller_id" not in h,
}

KIND_LABELS = {
    "genesys_calls": "Genesys Calls (CSV)",
    "genesys_status": "Genesys Status (CSV)",
    "csat": "CSAT (CSV)",
    "chat_timed_out": "Chat timed out (CSV)",
    "chat_closed": "Chat closed complete (CSV)",
    "incidents": "Incidents Resolved by SD (CSV)",
    "tasks": "Tasks Resolved by SD (CSV)",
    "rota": "Monthly ROTA (sheets L1 / L2 / BO)",
    "attendance": "Office attendance (DailyCheck table)",
    "team_fcr": "Team FCR (sheet 'L2 Tasks')",
    "quality": "Quality (sheet 'Merged')",
    "im": "Incident Management (sheet per month)",
    "km": "Knowledge Management",
}


def csv_header(path: Path) -> set[str]:
    with open(path, encoding="utf-8-sig", errors="replace", newline="") as f:
        return {h.strip() for h in next(csv.reader(f), [])}


def xlsx_kind(path: Path) -> str | None:
    wb = openpyxl.load_workbook(path, data_only=True)
    names = {ws.title.strip() for ws in wb.worksheets}
    tables = {t for ws in wb.worksheets for t in ws.tables.keys()}
    if {"L1", "L2", "BO"} <= names:
        return "rota"
    if "DailyCheck" in tables or "Daily Check" in names:
        return "attendance"
    if "L2 Tasks" in names:
        return "team_fcr"
    if "Merged" in names:
        return "quality"
    for ws in wb.worksheets:
        header = {str(c.value).strip() for c in ws[1] if c.value is not None}
        if "Response time Rate" in header:
            return "im"
        if "On-Time Delivery Rate" in header:
            return "km"
    return None


def discover(source: Path) -> dict[str, Path]:
    found: dict[str, list[Path]] = {}
    for path in sorted(source.iterdir()):
        if not path.is_file() or path.name.startswith(("~$", ".")):
            continue
        kind = None
        if path.suffix.lower() == ".csv":
            header = csv_header(path)
            kind = next((k for k, test in CSV_KINDS.items() if test(header)), None)
        elif path.suffix.lower() in (".xlsx", ".xlsm"):
            kind = xlsx_kind(path)
        if kind is None:
            log("INFO", f"Ignored (not a recognised source file): {path.name}")
            continue
        found.setdefault(kind, []).append(path)
    result = {}
    for kind, paths in found.items():
        if len(paths) > 1:
            raise SystemExit(
                f"ERROR: more than one {KIND_LABELS[kind]} file in '{source}':\n  - "
                + "\n  - ".join(p.name for p in paths)
                + "\nKeep only the file for the month you are reporting on."
            )
        result[kind] = paths[0]
        log("INFO", f"{KIND_LABELS[kind]:<42} <- {paths[0].name}")
    for kind, label in KIND_LABELS.items():
        if kind not in result:
            log("WARNING", f"No {label} file found - that part of the dashboard will be empty.")
    return result


# --------------------------------------------------------------------------- #
# Generic helpers
# --------------------------------------------------------------------------- #
def read_csv(path: Path) -> pd.DataFrame:
    df = pd.read_csv(path, dtype=str, keep_default_na=False, encoding="utf-8-sig")
    df.columns = [c.strip() for c in df.columns]
    return df


def parse_datetime(series: pd.Series) -> pd.Series:
    s = series.astype("string").str.strip().replace("", pd.NA)
    return pd.to_datetime(s, format="mixed", errors="coerce")


def in_month(dates: pd.Series, start: dt.date, end: dt.date) -> pd.Series:
    d = pd.to_datetime(dates, errors="coerce").dt.date
    return d.map(lambda x: x is not None and not pd.isna(x) and start <= x <= end).astype(bool)


def parse_clock(text):
    """' 00:05:51.940' -> (hours, minutes, seconds) or None."""
    if text is None or (isinstance(text, float) and pd.isna(text)):
        return None
    m = re.fullmatch(r"\s*(\d+):(\d{1,2}):(\d{1,2}(?:\.\d+)?)\s*", str(text))
    if not m:
        return None
    return int(m.group(1)), int(m.group(2)), float(m.group(3))


def clock_to_time(parts):
    """Power Query 'type time' (only valid below 24 hours)."""
    if parts is None or parts[0] >= 24:
        return None
    h, m, s = parts
    micro = int(round((s - int(s)) * 1_000_000))
    if micro >= 1_000_000:
        micro = 999_999
    return dt.time(h, m, int(s), micro)


def clock_to_timedelta(text, drop_fraction=False):
    parts = parse_clock(text)
    if parts is None:
        return pd.Timedelta(0)
    h, m, s = parts
    if drop_fraction:
        s = int(s)
    return pd.Timedelta(hours=h, minutes=m, seconds=s)


def hhmm(text: str) -> pd.Timedelta:
    h, m = text.split(":")
    return pd.Timedelta(hours=int(h), minutes=int(m))


def to_number(value):
    if value is None or (isinstance(value, float) and pd.isna(value)) or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def rate(value):
    """Rates are fractions (1 = 100 %). A value like 98.4 is converted to 0.984."""
    v = to_number(value)
    if v is not None and v > 1.5:
        v = v / 100
    return v


def sheet_rows(ws, header_row: int = 1):
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < header_row:
        return [], []
    header = rows[header_row - 1]
    return header, rows[header_row:]


def promote_header(values) -> list[str]:
    """Power Query PromoteHeaders: blank -> ColumnN, duplicates -> name_1."""
    out, seen = [], {}
    for i, v in enumerate(values, start=1):
        name = str(v).strip() if v is not None and str(v).strip() else f"Column{i}"
        if name in seen:
            seen[name] += 1
            name = f"{name}_{seen[name]}"
        else:
            seen[name] = 0
        out.append(name)
    return out


# --------------------------------------------------------------------------- #
# Table builders (each mirrors one query of the original report)
# --------------------------------------------------------------------------- #
def build_roster(path: Path | None, names: Names, settings: dict, start, end):
    """Agents and Schedule from the monthly ROTA (sheets L1, L2 and BO)."""
    agents, schedule = [], []
    if path is None:
        return pd.DataFrame(columns=["Agents", "Role"]), pd.DataFrame(columns=["Agents", "Date", "Shift"]), {}
    grade_to_role = {k.upper(): v for k, v in settings["grade_to_role"].items()}
    grades = {}
    wb = openpyxl.load_workbook(path, data_only=True)
    for sheet in ("L1", "L2", "BO"):
        ws = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))
        hdr_idx = next(
            (i for i, r in enumerate(rows) if len(r) > 1 and str(r[1] or "").strip().lower() == "grade"), None
        )
        if hdr_idx is None:
            log("WARNING", f"ROTA sheet '{sheet}': header row with 'Grade' not found - sheet skipped.")
            continue
        header = rows[hdr_idx]
        date_cols, started = [], False
        for ci in range(3, len(header)):
            v = header[ci]
            if isinstance(v, dt.datetime):
                started = True
                date_cols.append((ci, v.date()))
            elif started:
                break  # stop at the end of the calendar block (summary columns follow)
        seen_dates = set()
        date_cols = [(ci, d) for ci, d in date_cols if not (d in seen_dates or seen_dates.add(d))]
        for r in rows[hdr_idx + 1:]:
            raw = str(r[0]).strip() if r and r[0] is not None else ""
            if not raw or raw == "\xa0":
                break
            name = names.canon(raw)
            grade = str(r[1] or "").strip()
            grades[name] = grade
            role = grade_to_role.get(grade.upper())
            if role:
                agents.append((name, role))
            for ci, d in date_cols:
                if not (start <= d <= end) or ci >= len(r):
                    continue
                v = r[ci]
                shift = None
                if isinstance(v, dt.time):
                    shift = v
                elif isinstance(v, dt.datetime) and v.year <= 1900:
                    shift = v.time()
                elif isinstance(v, float) and 0 <= v < 1:
                    shift = (dt.datetime(1900, 1, 1) + dt.timedelta(days=v)).time()
                if shift is not None:
                    schedule.append((name, d, shift.replace(microsecond=0)))
    agents_df = pd.DataFrame(agents, columns=["Agents", "Role"]).drop_duplicates("Agents")
    sched_df = pd.DataFrame(schedule, columns=["Agents", "Date", "Shift"]).drop_duplicates(["Agents", "Date"])
    excluded = sorted(n for n, g in grades.items() if g.upper() not in grade_to_role)
    if excluded:
        log("INFO", f"ROTA people not shown on the scorecard (grade not in settings): {', '.join(excluded)}")
    names.set_roster(list(agents_df["Agents"]))
    return agents_df.sort_values(["Role", "Agents"]), sched_df.sort_values(["Agents", "Date"]), grades


GENESYS_COLUMNS = [
    "Conversation ID", "Date", "Agents", "Users", "Users - Interacted", "Users - Alerted",
    "Users - Not Responding", "Queue", "Duration", "Direction", "Abandoned", "Time to Abandon",
    "Total Queue", "Total Alert", "Total Talk", "Abandoned in Queue", "Disconnect Type", "Total Handle",
    "Total Queue Seconds", "Total Talk Hours", "Total Talk Minutes", "Total Talk Seconds",
    "Total Talk Calculated", "Total Queue Minutes",
]


def build_genesys(path, names: Names, settings, start, end):
    cols_ab = [c if c != "Users - Not Responding" else "Agentss" for c in GENESYS_COLUMNS]
    if path is None:
        return pd.DataFrame(columns=GENESYS_COLUMNS), pd.DataFrame(columns=cols_ab)
    df = read_csv(path)
    excluded = {_key(x) for x in settings["genesys_excluded_interacted_users"]}
    df = df[df["Conversation ID"].str.strip() != ""]
    df = df[(df["Users - Interacted"].str.strip() != "") & (~df["Users - Interacted"].map(_key).isin(excluded))]
    df = df[df["Direction"] == "Inbound"].copy()
    df["Date"] = parse_datetime(df["Date"])
    df = df[in_month(df["Date"], start, end)].copy()
    df["Date"] = df["Date"].dt.date
    df["Users - Interacted"] = df["Users - Interacted"].map(proper)
    df["Users - Not Responding"] = df["Users - Not Responding"].map(proper)

    queue = df["Total Queue"].map(parse_clock)
    talk = df["Total Talk"].map(parse_clock)
    df["Duration"] = df["Duration"].map(lambda x: clock_to_time(parse_clock(x)))
    df["Total Queue"] = queue.map(clock_to_time)
    df["Total Talk"] = talk.map(clock_to_time)
    q_ok = df["Total Queue"].notna()
    t_ok = df["Total Talk"].notna()
    df["Total Queue Seconds"] = [p[2] if ok else None for p, ok in zip(queue, q_ok)]
    df["Total Queue Minutes"] = [p[1] if ok else None for p, ok in zip(queue, q_ok)]
    df["Total Talk Hours"] = [p[0] if ok else None for p, ok in zip(talk, t_ok)]
    df["Total Talk Minutes"] = [p[1] if ok else None for p, ok in zip(talk, t_ok)]
    df["Total Talk Seconds"] = [p[2] if ok else None for p, ok in zip(talk, t_ok)]
    df["Total Talk Calculated"] = [p[0] * 60 + p[1] + p[2] / 60 if ok else None for p, ok in zip(talk, t_ok)]

    abandoned = df[df["Not Responding"].str.strip() != ""].copy()
    abandoned["Agentss"] = abandoned["Agents"].map(names.canon)
    abandoned["Agents"] = abandoned["Users - Not Responding"].map(names.canon)

    df["Agents"] = df["Agents"].map(names.canon)
    genesys = df[GENESYS_COLUMNS].sort_values(["Agents", "Date"])
    abandoned = abandoned[cols_ab].sort_values(["Agents", "Date"])
    return genesys, abandoned


def build_status_tables(path, names: Names, settings, start, end):
    """Status, Adherence (log in / log out per day) and Break, all from Genesys Status."""
    status_cols = ["Agents", "Date", "Start Time", "End Time", "Primary Status", "Duration",
                   "Duration String", "Total Hours", "Day Name"]
    adh_cols = ["Agents", "Date", "Log In Time", "Log Out Time", "Shift Duration", "Adherence"]
    brk_cols = ["Agents", "Date", "Primary Status", "Break Duration", "Break"]
    if path is None:
        return (pd.DataFrame(columns=status_cols), pd.DataFrame(columns=adh_cols), pd.DataFrame(columns=brk_cols))
    df = read_csv(path)
    df["Start Time"] = parse_datetime(df["Start Time"])
    df["End Time"] = parse_datetime(df["End Time"])
    df["Date"] = df["Start Time"].dt.date
    df = df[in_month(df["Start Time"], start, end)]
    df = df[~df["Primary Status"].isin(["Meal", "Offline"])].copy()
    df["Agents"] = df["Agent Name"].map(names.canon)
    df["Duration String"] = df["Duration"].map(lambda x: str(x).split(".")[0].strip())
    df["DurationTD"] = df["Duration"].map(lambda x: clock_to_timedelta(x, drop_fraction=True))
    df["Total Hours"] = df["DurationTD"].dt.total_seconds() / 3600
    df["Day Name"] = df["Start Time"].dt.day_name()

    status = df.assign(Duration=df["DurationTD"].dt.total_seconds() / 86400)[status_cols]
    status = status.sort_values(["Agents", "Start Time"])

    # Adherence: first log in / last log out per agent per day
    excluded_dates = {dt.date.fromisoformat(d) for d in settings.get("adherence_excluded_dates", [])}
    g = df.groupby(["Agents", "Date"], as_index=False).agg(LogIn=("Start Time", "min"), LogOut=("End Time", "max"))
    full, part = hhmm(settings["full_time_min_shift"]), hhmm(settings["part_time_min_shift"])
    part_time = {_key(a) for a in settings["part_time_agents"]}
    rows = []
    for r in g.itertuples(index=False):
        if r.Date in excluded_dates:
            continue
        shift = (r.LogOut - r.LogIn) if pd.notna(r.LogOut) else None
        ok = shift is not None and (shift >= full or (shift >= part and _key(r.Agents) in part_time))
        rows.append((
            r.Agents, r.Date, r.LogIn.time(),
            r.LogOut.time() if pd.notna(r.LogOut) else dt.time(11, 59),
            shift.total_seconds() / 86400 if shift is not None else None,
            "Positive" if ok else "Negative",
        ))
    adherence = pd.DataFrame(rows, columns=adh_cols).sort_values(["Agents", "Date"])

    # Break: total Away/Break time per agent per day
    b = df.assign(**{"Primary Status": df["Primary Status"].str.replace("Away", "Break")})
    b = b[b["Primary Status"] == "Break"]
    b = b.groupby(["Agents", "Date", "Primary Status"], as_index=False)["DurationTD"].sum()
    max_break = hhmm(settings["max_daily_break"])
    b["Break"] = b["DurationTD"].map(lambda x: "Positive" if x <= max_break else "Negative")
    b["Break Duration"] = b["DurationTD"].dt.total_seconds() / 3600
    breaks = b[brk_cols].sort_values(["Date", "Agents"], ascending=[False, True])
    return status, adherence, breaks


def build_csat(path, names: Names, settings, start, end):
    cols = ["Instance", "Instance.Task_ID", "Date", "Agents", "String Score", "Numerical Score", "Score"]
    if path is None:
        return pd.DataFrame(columns=cols)
    df = read_csv(path)
    df = df[df["instance.task_id.assigned_to"].str.strip() != ""].copy()
    df["Date"] = parse_datetime(df["instance.taken_on"])
    df = df[in_month(df["Date"], start, end)].copy()
    df["Date"] = df["Date"].dt.date
    scores = {"Very Satisfied": 5, "Satisfied": 4, "Neutral": 3, "Dissatisfied": 2, "Very Dissatisfied": 1}
    not_valid = set(settings["csat_not_valid_instances"])
    df["Numerical Score"] = df["string_value"].map(scores)

    def score(n, instance):
        if pd.isna(n):
            return None
        if n >= 3:
            return "Positive"
        return "Not Valid" if instance in not_valid else "Negative"

    df["Score"] = [score(n, i) for n, i in zip(df["Numerical Score"], df["instance"])]
    df["Agents"] = df["instance.task_id.assigned_to"].map(names.canon)
    df = df[df["Agents"].map(names.in_roster).astype(bool)]
    df = df.rename(columns={"instance": "Instance", "instance.task_id": "Instance.Task_ID", "string_value": "String Score"})
    return df[cols].sort_values(["Agents", "Date"])


def build_quality(path, names: Names, start, end):
    cols = ["Audit Id", "Contact Type", "Ticket Number", "Date", "Agents", "Total Score"]
    if path is None:
        return pd.DataFrame(columns=cols)
    df = pd.read_excel(path, sheet_name="Merged")
    df.columns = [str(c).strip() for c in df.columns]
    df["Date"] = pd.to_datetime(df["Audit Date"], errors="coerce")
    df = df[in_month(df["Date"], start, end)].copy()
    df["Date"] = df["Date"].dt.date
    df["Agents"] = df["Agent"].map(names.canon)
    df["Ticket Number"] = df["Ticket Number"].map(lambda x: None if pd.isna(x) else str(x).strip())
    df["Audit Id"] = pd.to_numeric(df["Audit Id"], errors="coerce").astype("Int64")
    df["Total Score"] = pd.to_numeric(df["Total Score"], errors="coerce")
    return df[cols].sort_values(["Agents", "Date"])


def build_team_fcr(path, names: Names, start):
    l2_cols = ["Agents", "Solved Assigned", "Total Assigned", "FCR Percentage", "Date"]
    team_cols = ["Team FCR", "Tasks", "Incidents", "Column4"]
    if path is None:
        return pd.DataFrame(columns=l2_cols), pd.DataFrame(columns=team_cols)
    wb = openpyxl.load_workbook(path, data_only=True)
    header, rows = sheet_rows(wb["L2 Tasks"])
    header = promote_header(header)
    l2 = pd.DataFrame([list(r) + [None] * (len(header) - len(r)) for r in rows], columns=header)
    l2 = l2[l2["Name"].notna() & (l2["Name"].astype(str).str.strip() != "")].copy()
    l2["Agents"] = l2["Name"].map(names.canon)
    l2["Solved Assigned"] = pd.to_numeric(l2["Solved Assigned"], errors="coerce").astype("Int64")
    l2["Total Assigned"] = pd.to_numeric(l2["Total Assigned"], errors="coerce").astype("Int64")
    l2["FCR Percentage"] = l2["FCR Percentage"].map(rate)
    l2["Date"] = start
    l2 = l2[l2_cols].drop_duplicates("Agents").sort_values("Agents")

    team = pd.DataFrame(columns=team_cols)
    if "All Team" in wb.sheetnames:
        header, rows = sheet_rows(wb["All Team"])
        header = promote_header(list(header)[:4])
        data = [list(r[:4]) for r in rows if any(v is not None for v in r[:4])]
        team = pd.DataFrame(data, columns=header).rename(columns=dict(zip(header, team_cols)))
        for c in ("Team FCR", "Tasks", "Incidents"):
            team[c] = team[c].map(lambda v: None if v is None else (str(int(v)) if isinstance(v, float) and v.is_integer() else str(v)))
        team["Column4"] = team["Column4"].map(to_number)
    else:
        log("WARNING", "Team FCR workbook has no 'All Team' sheet (only used by a hidden page).")
    return l2, team


MONTH_ABBR = {i: calendar.month_abbr[i].lower() for i in range(1, 13)}
MONTH_FULL = {i: calendar.month_name[i].lower() for i in range(1, 13)}


def month_sheet_matches(title: str, start: dt.date) -> bool:
    t = title.strip().lower()
    names = (MONTH_FULL[start.month], "sept" if start.month == 9 else None, MONTH_ABBR[start.month])
    if not any(n and re.match(rf"{n}\b", t) for n in names):
        return False
    years = re.findall(r"\d{2,4}", t)
    return not years or any(y in (str(start.year), str(start.year)[2:]) for y in years)


def latest_dates(files: dict[str, Path]) -> dict[str, dt.date]:
    """Last date found in each dated source file."""
    csv_date_cols = {
        "genesys_calls": "Date", "genesys_status": "Start Time", "csat": "instance.taken_on",
        "chat_closed": "opened_at", "chat_timed_out": "sys_created_on", "tasks": "opened_at", "incidents": "opened_at",
    }
    latest = {}
    for kind, column in csv_date_cols.items():
        if kind in files:
            d = parse_datetime(read_csv(files[kind])[column]).max()
            if pd.notna(d):
                latest[kind] = d.date()
    if "quality" in files:
        d = pd.to_datetime(pd.read_excel(files["quality"], sheet_name="Merged")["Audit Date"], errors="coerce").max()
        if pd.notna(d):
            latest["quality"] = d.date()
    if "attendance" in files:
        wb = openpyxl.load_workbook(files["attendance"], data_only=True)
        dates = [c.value for ws in wb.worksheets if "DailyCheck" in ws.tables or ws.title == "Daily Check"
                 for c in ws["A"] if isinstance(c.value, dt.datetime)]
        if dates:
            latest["attendance"] = max(dates).date()
    if "rota" in files:
        wb = openpyxl.load_workbook(files["rota"], data_only=True)
        rows = list(wb["L1"].iter_rows(min_row=1, max_row=5, values_only=True))
        dates = [v for r in rows for v in r if isinstance(v, dt.datetime)]
        if dates:
            # the ROTA calendar starts on a Monday, which can fall in the previous month
            months = pd.Series([d.date().replace(day=1) for d in dates]).mode()
            latest["rota"] = max(months)
    return latest


def detect_report_month(files: dict[str, Path]) -> dt.date:
    """Report month = the latest month that most source files reach.

    A single stray record (e.g. one CSAT survey on the 1st of the next month)
    does not move the dashboard to a new, almost empty month.
    """
    latest = latest_dates(files)
    if not latest:
        raise SystemExit("ERROR: no dated source files found - cannot work out the report month. Use --month YYYY-MM.")
    votes = pd.Series({k: d.replace(day=1) for k, d in latest.items()})
    counts = votes.value_counts()
    top = counts[counts == counts.max()].index.max()
    for kind, d in sorted(latest.items()):
        log("INFO", f"  last date in {KIND_LABELS[kind]:<42} {d}")
    others = sorted(k for k, m in votes.items() if m != top)
    if others:
        log("WARNING", "These files end in a different month than the rest: "
                       + ", ".join(f"{KIND_LABELS[k]} ({latest[k]:%d %b %Y})" for k in others))
    return top


def check_file_months(files: dict[str, Path], start: dt.date) -> None:
    """Warn when a file name mentions a different month (e.g. 'August ROTA.xlsx' for September)."""
    pattern = r"\b(" + "|".join(list(MONTH_FULL.values()) + [a for a in MONTH_ABBR.values()] + ["sept"]) + r")\b"
    for kind, path in files.items():
        found = re.findall(pattern, path.stem.lower())
        if not found:
            continue
        months = {next(i for i in range(1, 13) if MONTH_FULL[i].startswith(m[:3])) for m in found}
        if start.month not in months:
            log("WARNING", f"'{path.name}' looks like it is for another month than {start:%B %Y} - check it is the right file.")


def build_im(path, names: Names, agents: pd.DataFrame, start):
    metric_cols = ["Response time Rate", "Resolution time Rate", "Tasks on time Rate",
                   "First Reporting Time", "Reporting Accuracy Rate"]
    cols = ["Agents"] + metric_cols + ["Date"]
    if path is None:
        return pd.DataFrame(columns=cols)
    wb = openpyxl.load_workbook(path, data_only=True)

    def frame(ws):
        header, rows = sheet_rows(ws)
        header = promote_header(header)
        df = pd.DataFrame([list(r) + [None] * (len(header) - len(r)) for r in rows], columns=header)
        df = df[df["Name"].notna() & (df["Name"].astype(str).str.strip() != "")].copy()
        df["Agents"] = df["Name"].map(names.canon)
        return df

    matches = [ws for ws in wb.worksheets if month_sheet_matches(ws.title, start)]
    if matches:
        ws = matches[0]
        log("INFO", f"Incident Management: using sheet '{ws.title}'.")
    else:
        im_team = {_key(a) for a in agents.loc[agents["Role"] == "IM", "Agents"]}
        best = max(wb.worksheets, key=lambda w: len(im_team & set(frame(w)["Agents"].map(_key))))
        ws = best
        log("WARNING", f"Incident Management: no sheet for {start:%B %Y} "
                       f"(sheets: {', '.join(w.title for w in wb.worksheets)}). "
                       f"Using sheet '{ws.title}' - add a '{start:%b}' sheet to the workbook for correct figures.")
    df = frame(ws)
    for c in metric_cols:
        df[c] = df[c].map(rate) if c in df else None
    df["Date"] = start
    return df[cols].drop_duplicates("Agents").sort_values("Agents")


def build_km(path, names: Names, agents: pd.DataFrame):
    metric_cols = ["On-Time Delivery Rate", "Reporting Accuracy Rate", "First Reporting Time", "Reporting Accuracy Rate_1"]
    cols = metric_cols + ["Agents"]
    if path is None:
        return pd.DataFrame(columns=cols)
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = next(w for w in wb.worksheets if "On-Time Delivery Rate" in {str(c.value).strip() for c in w[1]})
    header, rows = sheet_rows(ws)
    header = promote_header(header)
    df = pd.DataFrame([list(r) + [None] * (len(header) - len(r)) for r in rows], columns=header)
    name_col = next((c for c in ("Agents", "Name") if c in df.columns), None)
    df = df[df[[c for c in metric_cols if c in df]].notna().any(axis=1)].copy()
    km_team = list(agents.loc[agents["Role"] == "KM", "Agents"])
    names_raw = df[name_col] if name_col else pd.Series([None] * len(df), index=df.index)
    if names_raw.isna().any() or (names_raw.astype(str).str.strip() == "").any():
        if len(km_team) == 1:
            log("WARNING", f"Knowledge Management: name cell is empty - figures assigned to {km_team[0]} "
                           f"(the only KM agent in the ROTA). Fill in the Name column to avoid this.")
            names_raw = names_raw.where(names_raw.notna() & (names_raw.astype(str).str.strip() != ""), km_team[0])
        else:
            log("WARNING", "Knowledge Management: rows without a name were ignored.")
    df["Agents"] = names_raw.map(names.canon)
    df = df[df["Agents"].notna() & (df["Agents"] != "")]
    for c in metric_cols:
        df[c] = df[c].map(rate) if c in df else None
    return df[cols].drop_duplicates("Agents")


def build_daily_check(path, names: Names, start, end):
    cols = ["Date", "Employee Name", "Required On-Site", "Manager Review", "Exception?", "Alert"]
    if path is None:
        return pd.DataFrame(columns=cols)
    wb = openpyxl.load_workbook(path, data_only=True)
    ws, ref = next(((ws, t.ref if hasattr(t, "ref") else t) for ws in wb.worksheets
                    for name, t in ws.tables.items() if name == "DailyCheck"), (None, None))
    if ws is None:
        ws = wb["Daily Check"]
        data = list(ws.iter_rows(values_only=True))
    else:
        data = [[c.value for c in row] for row in ws[ref]]
    header = promote_header(data[0])
    df = pd.DataFrame(data[1:], columns=header)
    df = df[df["Manager Review"].notna() & (df["Manager Review"].astype(str).str.strip() != "")].copy()
    df["Date"] = pd.to_datetime(df["Date"], errors="coerce")
    df = df[in_month(df["Date"], start, end)].copy()
    df["Date"] = df["Date"].dt.date
    df["Employee Name"] = df["Employee Name"].map(names.canon)
    for c in cols:
        if c not in df:
            df[c] = None
    return df[cols]


def build_chats(closed_path, timeout_path, names: Names, settings, start, end):
    cc_cols = ["number", "Date", "short_description", "opened_for", "type", "state", "Agents"]
    to_cols = ["Index", "Date", "work_item", "reason", "Agents"]
    closed = pd.DataFrame(columns=cc_cols)
    timeouts = pd.DataFrame(columns=to_cols)
    if closed_path is not None:
        df = read_csv(closed_path)
        df = df[df["state"] == "Closed Complete"].copy()
        df["Date"] = parse_datetime(df["opened_at"])
        df = df[in_month(df["Date"], start, end)].copy()
        df["Date"] = df["Date"].dt.date
        df["Agents"] = df["assigned_to"].map(names.canon)
        df = df.drop_duplicates("number")
        df = df[df["Agents"].map(names.in_roster).astype(bool)]
        closed = df[cc_cols]
    if timeout_path is not None:
        df = read_csv(timeout_path)
        df["Date"] = parse_datetime(df["sys_created_on"])
        df = df[in_month(df["Date"], start, end)].copy()
        df["Date"] = df["Date"].dt.date
        df["Agents"] = df["rejected_by"].map(names.canon)
        df = df[df["Agents"].map(names.in_roster).astype(bool)]
        df = df[~df["work_item"].isin(settings.get("chat_excluded_work_items", []))].copy()
        df.insert(0, "Index", range(1, len(df) + 1))
        timeouts = df[to_cols]
    return closed, timeouts


TASK_COLUMNS = ["number", "Date", "assigned_to", "state", "priority", "request_item.cat_item", "request_item",
                "request_item.u_request_on_behalf_of", "request_item.requested_for",
                "request_item.requested_for.location", "short_description", "assignment_group",
                "assignment_group.manager", "reassignment_count"]
INCIDENT_COLUMNS = ["number", "Date", "Agents", "short_description", "caller_id", "caller_id.location.country",
                    "priority", "state", "category", "subcategory", "business_service", "assignment_group",
                    "contact_type", "reassignment_count", "reopen_count", "close_code", "cmdb_ci", "u_rpt_mttr",
                    "opened_by", "incident_state", "request_item.cat_item", "request_item",
                    "request_item.u_request_on_behalf_of", "request_item.requested_for",
                    "request_item.requested_for.location", "assignment_group.manager"]


def build_resolved(path, columns, start, end, rename=None):
    if path is None:
        return pd.DataFrame(columns=columns)
    df = read_csv(path)
    if rename:
        df = df.rename(columns=rename)
    df["Date"] = parse_datetime(df["opened_at"])
    df = df[in_month(df["Date"], start, end)].copy()
    df["Date"] = df["Date"].dt.date
    for c in ("reassignment_count", "reopen_count", "u_rpt_mttr"):
        if c in df:
            df[c] = pd.to_numeric(df[c].replace("", pd.NA), errors="coerce").round().astype("Int64")
    for c in columns:
        if c not in df:
            df[c] = None
    return df[columns]


# --------------------------------------------------------------------------- #
# Output
# --------------------------------------------------------------------------- #
FORMATS = {dt.date: "yyyy-mm-dd", dt.time: "hh:mm:ss.000", dt.datetime: "yyyy-mm-dd hh:mm:ss"}


def cell_value(v):
    if v is None or v is pd.NA or v is pd.NaT:
        return None
    if isinstance(v, float) and pd.isna(v):
        return None
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime()
    if hasattr(v, "item") and not isinstance(v, (str, bytes)):
        try:
            return v.item()
        except (ValueError, AttributeError):
            pass
    if isinstance(v, str) and v == "":
        return None
    return v


def write_workbook(path: Path, tables: dict[str, pd.DataFrame]):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    bold = Font(bold=True)
    for name, df in tables.items():
        ws = wb.create_sheet(name)
        ws.append(list(df.columns))
        for c in ws[1]:
            c.font = bold
        for row in df.itertuples(index=False):
            ws.append([cell_value(v) for v in row])
        for col_cells in ws.iter_cols(min_row=2):
            for c in col_cells:
                fmt = FORMATS.get(type(c.value)) if c.value is not None else None
                if fmt:
                    c.number_format = fmt
        for i, col in enumerate(df.columns, start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(12, min(40, len(str(col)) + 4))
        ws.freeze_panes = "A2"
    tmp = path.with_name(path.stem + ".tmp.xlsx")
    wb.save(tmp)
    tmp.replace(path)


def main(argv=None):
    p = argparse.ArgumentParser(description="Prepare the DeLaval Scorecard data for one month.")
    p.add_argument("--month", help="Report month as YYYY-MM (default: the latest month found in the source files).")
    p.add_argument("--source", type=Path, default=DEFAULT_SOURCE, help="Folder with the raw source files.")
    p.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="Workbook to create.")
    args = p.parse_args(argv)

    for label, folder in (("Source files folder (DEFAULT_SOURCE)", args.source), ("Config folder (CONFIG_DIR)", CONFIG_DIR)):
        if not Path(folder).is_dir():
            raise SystemExit(f"ERROR: {label} not found: {folder}\n"
                             "Open scripts/prepare_scorecard_data.py and set the paths at the top of the file.")
    settings = load_settings()
    names = Names(load_aliases())
    files = discover(args.source.resolve())
    if args.month:
        start = dt.datetime.strptime(args.month, "%Y-%m").date().replace(day=1)
        log("INFO", f"Report month set with --month: {start:%B %Y}")
    else:
        start = detect_report_month(files)
    end = start.replace(day=calendar.monthrange(start.year, start.month)[1])
    log("INFO", f"Report month: {start:%B %Y} ({start} to {end})")
    check_file_months(files, start)
    f = files.get

    agents, schedule, _ = build_roster(f("rota"), names, settings, start, end)
    genesys, genesys_ab = build_genesys(f("genesys_calls"), names, settings, start, end)
    status, adherence, breaks = build_status_tables(f("genesys_status"), names, settings, start, end)
    csat = build_csat(f("csat"), names, settings, start, end)
    quality = build_quality(f("quality"), names, start, end)
    fcr_l2, fcr_team = build_team_fcr(f("team_fcr"), names, start)
    im = build_im(f("im"), names, agents, start)
    km = build_km(f("km"), names, agents)
    daily = build_daily_check(f("attendance"), names, start, end)
    chat_cc, chat_to = build_chats(f("chat_closed"), f("chat_timed_out"), names, settings, start, end)
    tasks = build_resolved(f("tasks"), TASK_COLUMNS, start, end)
    incidents = build_resolved(f("incidents"), INCIDENT_COLUMNS, start, end, rename={"assigned_to": "Agents"})

    report_month = pd.DataFrame([{
        "Month Start": start, "Month End": end, "Month Name": f"{start:%B %Y}",
        "Prepared At": dt.datetime.now().replace(microsecond=0),
    }])
    tables = {
        "ReportMonth": report_month,
        "Agents": agents,
        "Genesys": genesys,
        "Genesys_Abandoned": genesys_ab,
        "CSAT": csat,
        "Quality": quality,
        "Adherence": adherence,
        "Status": status,
        "Break": breaks,
        "Schedule": schedule,
        "FCR_L2": fcr_l2,
        "FCR_Team": fcr_team,
        "IM": im,
        "KM": km,
        "DailyCheck": daily,
        "Chat-Closed-Complete": chat_cc,
        "Chat-Time-Out": chat_to,
        "Tasks Resolved by SD": tasks,
        "Incidents Resolved by SD": incidents,
    }

    for name, df in tables.items():
        if name != "ReportMonth":
            log("INFO" if len(df) else "WARNING", f"{name:<26} {len(df):>6} rows")
    unmatched = sorted({a for df in (genesys, adherence, quality) for a in df["Agents"].dropna()
                        if a and not names.in_roster(a) and ";" not in a})
    if unmatched:
        log("INFO", "Names in the data that are not in the ROTA (not shown per agent; add them to "
                    "config/agent_aliases.csv if they are a spelling of a ROTA name): " + ", ".join(unmatched))

    tables["Log"] = pd.DataFrame(LOG, columns=["Level", "Message"])
    out = args.output.resolve()
    write_workbook(out, tables)
    print(f"\nDone. Workbook written to: {out}")
    print("Open the dashboard in Power BI Desktop and click Refresh.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
